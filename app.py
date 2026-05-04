import streamlit as st
from fpdf import FPDF
from fpdf.enums import XPos, YPos
from docx import Document
import data
import os
import io
import textwrap

# -------------------------------
# PAGE CONFIG
# -------------------------------
st.set_page_config(layout="wide")

BASE_DIR = os.path.dirname(os.path.abspath(__file__))

# -------------------------------
# TEXT CLEANER
# -------------------------------
def clean_text(text):
    if not isinstance(text, str):
        text = str(text)

    for k, v in {
        "–": "-",
        "—": "-",
        "’": "'",
        "“": '"',
        "”": '"',
        "•": "-",
        "✓": "-",
        "✅": "-"
    }.items():
        text = text.replace(k, v)

    return text.encode("latin-1", "ignore").decode("latin-1")

# -------------------------------
# UI STYLE
# -------------------------------
st.markdown("""
<style>
.stApp { background-color: #F4F8FB; max-width:1000px; margin:auto; }

.header-bar {
    background-color: #0072CE;
    padding: 20px;
    border-radius: 10px;
    color: white;
}

h2 {
    color: #0072CE;
    border-bottom: 2px solid #0072CE;
}

p, div {
    font-size: 15px;
    line-height: 1.6;
}
</style>
""", unsafe_allow_html=True)

# -------------------------------
# HEADER
# -------------------------------
st.markdown("""
<div class="header-bar">
<h1>Raj Singh</h1>
<p><b>Senior HSE Manager</b></p>
<p>Raj.singh.hse@gmail.com | +91 9883122847</p>
</div>
""", unsafe_allow_html=True)

# IMAGE
img_path = os.path.join(BASE_DIR, "profile.jpg")
if os.path.exists(img_path):
    st.image(img_path, width=180)

# -------------------------------
# CONTENT
# -------------------------------
st.header("Skills")
for s in data.skills:
    st.markdown(f"- {s}")

st.header("Professional Experience")
for exp in data.experience:
    st.subheader(f"{exp['role']} — {exp['company']}")
    for d in exp["details"]:
        st.markdown(f"- {d}")

st.header("Key Achievements")
for c, items in data.achievements.items():
    st.subheader(c)
    for i in items:
        st.markdown(f"- {i}")

st.header("Education")
for e in data.education:
    st.markdown(f"- {e}")

st.header("Certifications")
for c in data.certifications:
    st.markdown(f"- {c}")

# -------------------------------
# ✅ FINAL FIXED PDF FUNCTION
# -------------------------------
def create_pdf():
    pdf = FPDF()
    pdf.add_page()
    pdf.set_auto_page_break(True, 10)

    # FIXED margins
    pdf.set_left_margin(10)
    pdf.set_right_margin(10)

    BLUE = (0, 114, 206)
    BLACK = (0, 0, 0)

    # Full usable width
    page_width = pdf.w - pdf.l_margin - pdf.r_margin

    # HEADER BAR
    pdf.set_fill_color(*BLUE)
    pdf.rect(0, 0, pdf.w, 25, style='F')

    pdf.set_text_color(255, 255, 255)
    pdf.set_font("Arial", "B", 16)
    pdf.set_xy(10, 10)
    pdf.cell(0, 5, "Raj Singh - Senior HSE Manager")

    pdf.set_font("Arial", "", 10)
    pdf.set_xy(10, 18)
    pdf.cell(0, 5, "Raj.singh.hse@gmail.com | +91 9883122847")

    pdf.ln(20)

    # ✅ SECTION FUNCTION (ALIGNMENT FIXED)
    def section(title, items):
        pdf.set_text_color(*BLUE)
        pdf.set_font("Arial", "B", 13)

        pdf.set_x(pdf.l_margin)
        pdf.multi_cell(page_width, 8, clean_text(title))

        pdf.set_draw_color(*BLUE)
        pdf.line(pdf.l_margin, pdf.get_y(), pdf.w - pdf.r_margin, pdf.get_y())

        pdf.ln(3)

        pdf.set_text_color(*BLACK)
        pdf.set_font("Arial", "", 11)

        for item in items:
            pdf.set_x(pdf.l_margin)  # ✅ force left alignment
            text = clean_text(item)

            wrapped = textwrap.wrap(text, 90)

            for line in wrapped:
                pdf.multi_cell(page_width, 6, "- " + line)

        pdf.ln(2)

    # CALLS
    section("Skills", data.skills)
    section("Education", data.education)
    section("Certifications", data.certifications)

    # ✅ EXPERIENCE FIXED
    pdf.set_text_color(*BLUE)
    pdf.set_font("Arial", "B", 13)

    pdf.set_x(pdf.l_margin)
    pdf.multi_cell(page_width, 8, "Experience")

    pdf.line(pdf.l_margin, pdf.get_y(), pdf.w - pdf.r_margin, pdf.get_y())
    pdf.ln(3)

    pdf.set_text_color(*BLACK)

    for exp in data.experience:
        pdf.set_font("Arial", "B", 11)

        pdf.set_x(pdf.l_margin)
        header = clean_text(f"{exp.get('role')} - {exp.get('company')}")
        pdf.multi_cell(page_width, 7, header)

        pdf.set_font("Arial", "", 11)

        for d in exp.get("details", []):
            pdf.set_x(pdf.l_margin)
            text = clean_text(d)

            wrapped = textwrap.wrap(text, 90)

            for line in wrapped:
                pdf.multi_cell(page_width, 6, "- " + line)

        pdf.ln(2)

    # ✅ ACHIEVEMENTS
    section("Achievements", [
        f"{c}: {i}" for c, items in data.achievements.items() for i in items
    ])

    return bytes(pdf.output())

# -------------------------------
# WORD EXPORT
# -------------------------------
def create_word():
    doc = Document()

    doc.add_heading("Raj Singh - Senior HSE Manager", 0)
    doc.add_paragraph("Raj.singh.hse@gmail.com | +91 9883122847")

    for section, content in {
        "Skills": data.skills,
        "Education": data.education,
        "Certifications": data.certifications
    }.items():
        doc.add_heading(section, 1)
        for item in content:
            doc.add_paragraph(str(item), style="List Bullet")

    doc.add_heading("Experience", 1)
    for exp in data.experience:
        doc.add_paragraph(f"{exp['role']} - {exp['company']}", style="List Bullet")
        for d in exp["details"]:
            doc.add_paragraph(str(d), style="List Bullet 2")

    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()

# -------------------------------
# DOWNLOAD
# -------------------------------
st.write("---")

pdf_data = create_pdf()
st.download_button("📄 Download Premium PDF", pdf_data, "Raj_Singh_Premium.pdf")

word_data = create_word()
st.download_button("📝 Download Word", word_data, "Raj_Singh_Profile.docx")